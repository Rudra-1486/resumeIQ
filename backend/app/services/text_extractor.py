"""In-memory text extractor and whitespace normalizer for PDF and DOCX files.

PRD References:
- Section 3: Supported Document Types & Limitations (No OCR, 50 word minimum)
- Section 14: AI Analysis Workflow (Step 3: In-Memory Text Extraction)
- Section 18.3: Endpoint 2: Analyze Resume (File constraints & metadata)
- Section 21.2: Volatile Memory Processing Only (Zero disk I/O)
- Section 22: Error Handling (EC-05, EC-06, EC-07, EC-08, EC-09)
"""

import io
import math
import re
from dataclasses import dataclass
from typing import List, Optional

try:
    import pymupdf as fitz
except ImportError:
    import fitz

import docx
from docx import Document

from app.config import MAX_EXTRACTED_CHARS, MIN_EXTRACTED_WORDS
from app.utils.exceptions import (
    CorruptedDocumentError,
    EmptyDocumentError,
    UnreadableOrScannedDocumentError,
)

MAX_PDF_PAGES: int = 10
MIN_NON_WHITESPACE_CHARS: int = 250


@dataclass
class ExtractedDocument:
    """Standardized representation of in-memory extracted document content."""

    text: str
    word_count: int
    char_count: int
    page_count: int
    file_type: str
    text_truncated: bool


def normalize_whitespace(text: str) -> str:
    """Normalize whitespace while preserving line breaks, paragraph structure, and bullets.

    - Replaces non-breaking space (\\u00a0) with standard space.
    - Strips trailing whitespace per line.
    - Collapses 3 or more consecutive newlines into 2 newlines.
    """
    if not text:
        return ""

    # Replace non-breaking spaces and zero-width spaces
    text = text.replace("\u00a0", " ").replace("\ufeff", "")

    # Strip trailing whitespace on each line
    lines = [line.rstrip() for line in text.splitlines()]
    clean_text = "\n".join(lines)

    # Collapse excessive blank lines
    clean_text = re.sub(r"\n{3,}", "\n\n", clean_text)
    return clean_text.strip()


def extract_text_from_pdf(file_bytes: bytes) -> ExtractedDocument:
    """Parse text from PDF bytes entirely in volatile memory via PyMuPDF.

    Args:
        file_bytes: Raw binary content of the PDF.

    Returns:
        ExtractedDocument: Extracted and normalized text model.

    Raises:
        CorruptedDocumentError: If the PDF is encrypted, password-protected, or corrupted.
        EmptyDocumentError: If document contains 0 words/characters.
        UnreadableOrScannedDocumentError: If document contains < 50 words or < 250 non-ws chars.
    """
    try:
        doc = fitz.open(stream=file_bytes, filetype="pdf")
    except Exception as exc:
        raise CorruptedDocumentError(
            message="This PDF could not be opened because it is password-protected or corrupted. Please upload an unencrypted PDF.",
            code="CORRUPTED_OR_ENCRYPTED_PDF",
            details={"error_reason": str(exc)},
        ) from exc

    try:
        if doc.is_encrypted:
            raise CorruptedDocumentError(
                message="This PDF could not be opened because it is password-protected or corrupted. Please upload an unencrypted PDF.",
                code="CORRUPTED_OR_ENCRYPTED_PDF",
                details={"is_encrypted": True},
            )

        total_pages = doc.page_count
        pages_to_extract = min(total_pages, MAX_PDF_PAGES)
        text_truncated = total_pages > MAX_PDF_PAGES

        extracted_lines: List[str] = []
        for page_idx in range(pages_to_extract):
            page = doc[page_idx]
            # Use sort=True for natural reading order across columns/blocks
            blocks = page.get_text("blocks", sort=True)
            for block in blocks:
                # block: (x0, y0, x1, y1, text, block_no, block_type)
                # block_type == 0 denotes text block (1 is image)
                if len(block) > 6 and block[6] == 0:
                    b_text = block[4].strip()
                    if b_text:
                        extracted_lines.append(b_text)
                elif len(block) > 4 and isinstance(block[4], str):
                    b_text = block[4].strip()
                    if b_text:
                        extracted_lines.append(b_text)

        raw_text = "\n\n".join(extracted_lines)
    finally:
        doc.close()

    normalized = normalize_whitespace(raw_text)

    if len(normalized) > MAX_EXTRACTED_CHARS:
        normalized = normalized[:MAX_EXTRACTED_CHARS]
        text_truncated = True

    words = normalized.split()
    word_count = len(words)
    non_ws_count = len(re.sub(r"\s+", "", normalized))

    if word_count == 0 or non_ws_count == 0:
        raise EmptyDocumentError(
            "The selected document appears to be empty. Please upload a resume containing readable text."
        )

    if word_count < MIN_EXTRACTED_WORDS or non_ws_count < MIN_NON_WHITESPACE_CHARS:
        raise UnreadableOrScannedDocumentError(
            extracted_word_count=word_count,
            minimum_required_words=MIN_EXTRACTED_WORDS,
        )

    return ExtractedDocument(
        text=normalized,
        word_count=word_count,
        char_count=len(normalized),
        page_count=total_pages,
        file_type="pdf",
        text_truncated=text_truncated,
    )


def extract_text_from_docx(file_bytes: bytes) -> ExtractedDocument:
    """Parse text from DOCX bytes entirely in volatile memory via python-docx.

    Extracts paragraphs and table cells (deduplicating adjacent identical cell strings).

    Args:
        file_bytes: Raw binary content of the DOCX file.

    Returns:
        ExtractedDocument: Extracted and normalized text model.

    Raises:
        CorruptedDocumentError: If the Word archive is corrupted.
        EmptyDocumentError: If document contains 0 words/characters.
        UnreadableOrScannedDocumentError: If document contains < 50 words or < 250 non-ws chars.
    """
    try:
        doc = Document(io.BytesIO(file_bytes))
    except Exception as exc:
        raise CorruptedDocumentError(
            message="This Word document appears to be corrupted and could not be read. Please re-save it as a .docx file and try again.",
            code="CORRUPTED_DOCX",
            details={"error_reason": str(exc)},
        ) from exc

    extracted_lines: List[str] = []

    # 1. Extract paragraphs
    for para in doc.paragraphs:
        p_text = para.text.strip()
        if p_text:
            extracted_lines.append(p_text)

    # 2. Extract table cells (with deduplication of adjacent identical strings from merged cells)
    for table in doc.tables:
        for row in table.rows:
            row_cells: List[str] = []
            last_cell_text: Optional[str] = None
            for cell in row.cells:
                c_text = cell.text.strip()
                if c_text and c_text != last_cell_text:
                    row_cells.append(c_text)
                    last_cell_text = c_text
            if row_cells:
                extracted_lines.append(" | ".join(row_cells))

    raw_text = "\n".join(extracted_lines)
    normalized = normalize_whitespace(raw_text)

    text_truncated = False
    if len(normalized) > MAX_EXTRACTED_CHARS:
        normalized = normalized[:MAX_EXTRACTED_CHARS]
        text_truncated = True

    words = normalized.split()
    word_count = len(words)
    non_ws_count = len(re.sub(r"\s+", "", normalized))

    if word_count == 0 or non_ws_count == 0:
        raise EmptyDocumentError(
            "The selected document appears to be empty. Please upload a resume containing readable text."
        )

    if word_count < MIN_EXTRACTED_WORDS or non_ws_count < MIN_NON_WHITESPACE_CHARS:
        raise UnreadableOrScannedDocumentError(
            extracted_word_count=word_count,
            minimum_required_words=MIN_EXTRACTED_WORDS,
        )

    # Word documents don't store strict page boundaries in XML; estimate ~450 words/page
    estimated_pages = max(1, math.ceil(word_count / 450))

    return ExtractedDocument(
        text=normalized,
        word_count=word_count,
        char_count=len(normalized),
        page_count=estimated_pages,
        file_type="docx",
        text_truncated=text_truncated,
    )


def extract_text_from_bytes(file_bytes: bytes, file_type: str) -> ExtractedDocument:
    """Route document parsing to PDF or DOCX extractor in memory.

    Args:
        file_bytes: Raw binary content of the file.
        file_type: Document format string ("pdf" or "docx").

    Returns:
        ExtractedDocument: Dataclass containing extracted text, counts, and metadata.
    """
    normalized_type = file_type.lower().strip().replace(".", "")
    if normalized_type == "pdf":
        return extract_text_from_pdf(file_bytes)
    elif normalized_type == "docx":
        return extract_text_from_docx(file_bytes)
    else:
        raise CorruptedDocumentError(
            f"Unsupported file format '{file_type}' for extraction.",
            code="INVALID_FILE_TYPE",
        )


# Alias for backward compatibility and alternate naming
extract_document_text = extract_text_from_bytes

